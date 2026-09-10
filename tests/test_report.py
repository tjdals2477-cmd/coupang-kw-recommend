import io
import unittest

from openpyxl import load_workbook

from coupang_kw_rec.config import load_config
from coupang_kw_rec.filter import FilterResult
from coupang_kw_rec.naver_api import CollectionResult
from coupang_kw_rec.pipeline import PipelineResult
from coupang_kw_rec.report import build_workbook_bytes
from coupang_kw_rec.seed import SeedChoice


class ReportTests(unittest.TestCase):
    def test_sheet_and_recommendation_column_contract(self):
        config = load_config()
        row = {
            "rank": 1, "keyword": "LED일자등", "rec_grade": "REC_A_즉시등록",
            "action_ko": "광고센터 등록", "score": 0.9, "search_volume": 100,
            "pc_qc": 40, "mobile_qc": 60, "mobile_ratio": 0.6, "comp_idx": "중간",
            "pl_avg_depth": 5, "naver_avg_ctr": 0.02, "seed_keyword": "LED등",
            "seed_count": 1, "seed_roas": float("nan"), "relevance": 1.0,
            "brand": "NONE", "already_converting": False, "note": "",
        }
        result = PipelineResult(
            [SeedChoice("LED등", "직접입력", 1)], CollectionResult(total_chunks=1, successful_chunks=1),
            FilterResult(candidates=[row]), [row], [row], 0, [],
        )
        workbook = load_workbook(io.BytesIO(build_workbook_bytes(result, config)), read_only=True)
        self.assertEqual(workbook.sheetnames, list(config["output"]["sheets"].values()))
        headers = [cell.value for cell in next(workbook[config["output"]["sheets"]["recommendations"]].iter_rows(max_row=1))]
        self.assertEqual(headers, config["output"]["recommendation_columns"])


if __name__ == "__main__":
    unittest.main()
